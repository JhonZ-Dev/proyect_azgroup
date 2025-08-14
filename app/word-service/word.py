from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn

def set_cell_bg(cell, color_hex):
    """Set background color of a cell by direct XML injection."""
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{}"/>'.format(color_hex)
        )
    )

def generar_doc_proforma(data, plantilla_path="plantilla.docx", archivo_salida="proforma_final.docx"):
    # 1. Cargar plantilla membretada y eliminar tabla dummy
    doc = Document(plantilla_path)
    if doc.tables:
        tbl = doc.tables[0]
        tbl._element.getparent().remove(tbl._element)

    # 2. Crear tabla general (7 columnas)
    tabla = doc.add_table(rows=0, cols=7)
    tabla.style = 'Table Grid'

    # --- PROFORMA (cabecera)
    row = tabla.add_row().cells
    row[0].merge(row[6])
    p = row[0].paragraphs[0]
    run = p.add_run(f"PROFORMA {data.get('txt_necesidad', '')}")
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = 'Arial'
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_bg(row[0], 'dae9f7')

    # --- Campos principales (campo + respuesta)
    campos = [
        ('CLIENTE', data.get('txt_cliente', '')),
        ('RUC', data.get('txt_ruc', '')),
        ('DIRECCIÓN:', data.get('txt_direccion', '')),
        ('FECHA', data.get('txt_fecha', '')),
        ('TELEFONO', data.get('txt_telefono', '')),
        ('NECESIDAD', data.get('txt_necesidad', '')),
        ('FUNCIONARIO ENCARGADO:', data.get('txt_funcionario', '')),
        ('CORREO', data.get('txt_correo', '')),
        ('HORA MAXIMA:', str(data.get('tHora_maxina', ''))[:16].replace('T',' ')),
    ]
    for label, valor in campos:
        row = tabla.add_row().cells
        row[0].merge(row[1])
        row[0].text = label
        row[2].merge(row[6])
        row[2].text = valor

    # --- OBJETO DE COMPRA (título y respuesta)
    row = tabla.add_row().cells
    row[0].merge(row[6])
    row[0].text = "OBJETO DE COMPRA"
    row[0].paragraphs[0].runs[0].bold = True

    row = tabla.add_row().cells
    row[0].merge(row[6])
    row[0].text = data.get('txt_objetivoCompra', '')

    # --- CABECERAS DE ÍTEMS
    headers = ['No.', 'CPC', 'UNIDAD', 'ESPECIFICACIONES', 'CANTIDAD', 'P. UNIT', 'P.TOTAL']
    row = tabla.add_row().cells
    for i, h in enumerate(headers):
        row[i].text = h
        row[i].paragraphs[0].runs[0].bold = True

    # --- FILAS DE ÍTEMS
    items = data.get('items', [])
    total = 0
    for idx, item in enumerate(items):
        row = tabla.add_row().cells
        row[0].text = str(idx+1)
        row[1].text = str(item.get('txt_cpc', ''))
        row[2].text = str(item.get('txt_unidad', ''))
        row[3].text = str(item.get('txt_especificaciones', ''))
        row[4].text = str(item.get('int_cantidad', ''))
        row[5].text = f"{item.get('flo_precioUnitario', 0):.2f}"
        row[6].text = f"{item.get('flo_precioTotal', 0):.2f}"
        total += item.get('flo_precioTotal', 0)

    # --- TOTAL
    row = tabla.add_row().cells
    row[0].merge(row[5])
    row[0].text = "TOTAL"
    row[6].text = f"${total:.2f}"

    # --- NO GRAVAMOS IVA
    row = tabla.add_row().cells
    row[0].merge(row[6])
    row[0].text = "NO GRAVAMOS IVA - SOMOS REGIMEN RIMPE - NEGOCIO POPULAR"

    # --- CAMPOS FINALES
    finales = [
        ('PLAZO DE ENTREGA', data.get('txt_plazoEntrega', '')),
        ('VIGENCIA DE LA OFERTA', data.get('txt_vigenciaOferta', '')),
        ('FORMA DE PAGO:', data.get('txt_formaPago', '')),
        ('METODOLOGIA DE TRABAJO', data.get('txt_metodologiaTrabajo', '')),
        ('GARANTIA', data.get('txt_garantia', '')),
        ('ENLACE', data.get('txt_enlace', ''))
    ]
    for label, valor in finales:
        row = tabla.add_row().cells
        row[0].merge(row[1])
        row[0].text = label
        row[2].merge(row[6])
        row[2].text = valor

    # --- Firma fuera de la tabla
    doc.add_paragraph('')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('DAYANA LISBETH ZAMBRANO MACIAS\nREPRESENTANTE LEGAL\nRUC: 2350621211001')
    run.bold = True
    run.font.size = Pt(14)

    doc.save(archivo_salida)
    return archivo_salida